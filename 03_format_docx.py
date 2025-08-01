# -*- coding: utf-8 -*-
# ---------------------------------------------------------------------------
# Format Taxa Descriptions
# Author: Amanda Droghini
# Last Updated: 2025-07-30
# Usage: Execute in Python 3.13+.
# Description: "Format Taxa Descriptions" reads in Word documents, separates section headings from text, re-formats
# the document using Markdown, and adds Hugo front matter. The output is a Markdown file that can be converted to a
# HTML page using the static website generator Hugo.
# ---------------------------------------------------------------------------

# Import required libraries
from pathlib import Path
import re
import textwrap
from docx import Document
import polars as pl
from utils import collect_docx_info

# Set root directory
drive = Path('C:/')
root_folder = 'ACCS_Work'

# Define folder structure
project_folder = drive / root_folder / 'Projects' / 'Lichen_Guide'
taxa_folder = project_folder / 'Guide Master Folder_V_7_16_25' / 'Taxa Folders'
website_folder = drive / root_folder / 'Servers_Websites' / 'akveg-lichens'

# Define output folder
output_folder = website_folder / 'content' / 'pages' / 'taxa'

# Define input files
hierarchy_input = project_folder / 'outputs' / 'taxon_hierarchy.csv'
thumbnail_input = project_folder / 'outputs' / 'thumbnail_files.csv'

# Ingest input files
hierarchy = pl.read_csv(hierarchy_input)
thumbnails = pl.read_csv(thumbnail_input, columns = ['taxon_name', 'short_code',
                                                     'sequence_number', 'output_name', 'thumbnail_path'],
                         schema_overrides={'sequence_number':pl.Int64})

# Identify all .docx files in subdirectories of taxa_folder
docx_list = collect_docx_info(taxa_folder)
docx_files = pl.DataFrame(docx_list)

# Obtain name of output folder based on taxon organization
docx_hierarchy = docx_files.join(hierarchy, left_on='taxon_name', right_on='original_folder', how='left')

## Fill in non-matches (null values)
docx_hierarchy = docx_hierarchy.with_columns(
    pl.when((pl.col("taxon_folder").is_null()) & pl.col('taxon_name').str.contains("Cladonia"))
    .then(pl.lit("cladoniaceae"))
    .when((pl.col("taxon_folder").is_null()) & pl.col('taxon_name').str.contains("Thamnolia|Siphula|Lepra"))
    .then(pl.lit("icmadophilaceae"))
    .when((pl.col("taxon_folder").is_null()) & pl.col('taxon_name').str.contains("Dactylina|Allocetraria"))
    .then(pl.lit("non_shrub_hair"))
    .when(pl.col("taxon_folder").is_null())
    .then(pl.col("taxon_name"))
    .otherwise(pl.col("taxon_folder"))
    .alias("taxon_folder")
)

## Replace remaining missing matches
text_mapping = {"Glypholecia scabra":"scales_squamule_like",
                "Lobaria linita & Lobaria tenuior": "lungworts",
                "Rusavskia elegans & Rusavskia sorediata": "teloschistaceae",
                "Sporastatia polyspora & Sporastatia testudinea": "crusts_fruticose","Xanthoparmelia spp": "shield_like_parmelioids",
}

docx_hierarchy = docx_hierarchy.with_columns(
    pl.col("taxon_folder").replace(text_mapping)  # Keeps original value if not listed in dictionary
    .alias("taxon_folder")
)

## Ensure that all null values have been addressed
print(docx_hierarchy['taxon_folder'].is_null().sum())

# Define output paths based on taxon organization
docx_hierarchy = (docx_hierarchy.with_columns(
    pl.concat_str(str(output_folder) +
                  pl.lit("\\") +
                  pl.col("taxon_folder") +
                  pl.lit("\\") +
                  pl.col("short_code") +
                  pl.lit(".md"))
    .alias('output_path')
)
                  .select(['taxon_name', 'input_docx', 'output_path']))

# Process taxonomic description into Markdown file

# Generate temp to test
temp = docx_hierarchy.filter(pl.col('taxon_name').str.contains("Cladonia"))[0:10,:]

for row in temp.iter_rows(named=True):
    docx_path = row['input_docx']  # Keep as string
    md_path = Path(row['output_path'])
    taxon_name = row['taxon_name']

    print(f"Formatting Markdown for {taxon_name}")

    # List image file names associated with taxon
    taxon_images = thumbnails.filter(
        pl.col('taxon_name') == taxon_name)['output_name'].to_list()

    # Open Word document
    document = Document(docx_path)

    # Create blank list to store formatted outputs
    markdown_parts = []

    # Write front matter
    front_matter = textwrap.dedent(f"""\
        ---
        title: "{taxon_name}"
        type: docs
        ---
        """)

    markdown_parts.append(front_matter)

    # Write Heading 1
    first_heading = f"# {taxon_name}"

    markdown_parts.append(first_heading)

    # Insert first image (if present)
    if len(taxon_images) > 0:
        first_img = taxon_images[0]
        first_img_path = f"/images/taxa/{first_img}"
        first_img_txt = f"![{taxon_name}]({first_img_path})"
        markdown_parts.append(first_img_txt)

    # Format subsequent headings + paragraph text
    for para in document.paragraphs:
        split_para = re.split(pattern=r':\s(.*)', string=para.text, maxsplit=2)

        if len(split_para) > 1:
            para_heading = split_para[0].strip()
            para_text = split_para[1].strip()
            if para_heading != "Name" and para_text:
                # Format heading
                formatted_heading = f"## {para_heading}"
                # Concatenate heading and paragraph text
                formatted_line = formatted_heading + "\n" + para_text
                formatted_line = formatted_line.rstrip('\n')
                markdown_parts.append(formatted_line)

        elif len(para.text) < 15:  # Consider only excluding empty strings
            print(f"Skipping writing string: ", para.text)

        else:
            formatted_line = para.text.rstrip('\n')
            markdown_parts.append(formatted_line)

    # Append other images (if present)
    if len(taxon_images) > 1:
        photo_heading = "## Photos"
        markdown_parts.append(photo_heading)

        for img_file in taxon_images[1:]:
            img_path = f"/images/taxa/{img_file}"
            img_txt = f"![{taxon_name}]({img_path})"
            markdown_parts.append(img_txt)

    # Combine Markdown parts, separating each with two newlines
    compiled_markdown = "\n\n".join(markdown_parts)

    # Remove all newlines at the end of the document and replace with a single one
    final_markdown = compiled_markdown.rstrip('\n') + '\n'

    # Write to disk
    with (open(md_path, 'w', encoding='utf-8') as md_file):
        md_file.write(final_markdown)
